"""Экспорт отчётов в XLSX (openpyxl) и PDF (reportlab)."""

import io
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from apps.reports.reports import Report

FONT_DIR = Path(__file__).resolve().parent.parent.parent / "static" / "fonts"

_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_TOTAL_FONT = Font(bold=True)


def _excel_lines(spec: Report, rows):
    yield list(spec.labels)
    for row in rows:
        yield ["" if row.get(key) is None else row.get(key) for key in spec.keys]


def write_xlsx_bytes(spec: Report, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = f"Приложение №{spec.number}"
    for i, line in enumerate(_excel_lines(spec, rows), start=1):
        for col, value in enumerate(line, start=1):
            c = ws.cell(row=i, column=col, value=value)
            if i == 1:
                c.fill = _HEADER_FILL
                c.font = _HEADER_FONT
                c.alignment = Alignment(horizontal="center")
            elif line and str(line[0]).startswith("ИТОГО"):
                c.font = _TOTAL_FONT
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col) + 2
        ws.column_dimensions[col[0].column_letter].width = min(width, 40)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def write_pdf(spec: Report, rows):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    if "DejaVuSans" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("DejaVuSans", str(FONT_DIR / "DejaVuSans.ttf")))
        pdfmetrics.registerFont(
            TTFont("DejaVuSans-Bold", str(FONT_DIR / "DejaVuSans-Bold.ttf"))
        )

    wide = len(spec.columns) > 8
    page_size = landscape(A4) if wide else A4
    title_style = ParagraphStyle(
        "Title",
        fontName="DejaVuSans-Bold",
        fontSize=13,
        leading=17,
        spaceAfter=6,
    )
    meta_style = ParagraphStyle("Meta", fontName="DejaVuSans", fontSize=8, leading=10)
    cell_style = ParagraphStyle(
        "Cell", fontName="DejaVuSans", fontSize=7, leading=9
    )
    head_style = ParagraphStyle(
        "Head",
        fontName="DejaVuSans-Bold",
        fontSize=7,
        leading=9,
        textColor=colors.white,
    )

    data = [[Paragraph(label, head_style) for label in spec.labels]]
    for row in rows:
        total = bool(row.get("_total"))
        cells = []
        for key in spec.keys:
            value = row.get(key)
            value = "" if value is None else str(value)
            style = cell_style
            if total:
                style = ParagraphStyle("T", parent=cell_style, fontName="DejaVuSans-Bold")
            cells.append(Paragraph(value, style))
        data.append(cells)

    width, height = page_size
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=page_size,
        leftMargin=20,
        rightMargin=20,
        topMargin=24,
        bottomMargin=24,
        title=f"Отчёт: {spec.title}",
    )
    story = [
        Paragraph(spec.title, title_style),
        Paragraph(
            f"Приложение №{spec.number} к ТЗ · еж. — {spec.title}",
            meta_style,
        ),
        Spacer(1, 8),
        table,
    ]
    doc.build(story)
    return buf.getvalue()
