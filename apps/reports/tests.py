"""Тесты модуля отчётов: реестр, фильтры, расчётные формы, экспорт (Этап 5)."""

import datetime
import io
import uuid

from django.contrib.auth.models import Group
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import load_workbook

from apps.core.roles import ensure_role_groups
from apps.employee.models import Employee
from apps.journal.models import Irp, IrpTheme
from apps.reports.export import write_pdf, write_xlsx_bytes
from apps.reports.forms import ReportFilterForm
from apps.reports.reports import REPORT_INDEX, REPORTS, ReportFilters


class ReportRegistryTests(TestCase):
    def test_nine_reports(self):
        self.assertEqual(len(REPORTS), 9)

    def test_slugs_unique(self):
        slugs = [r.slug for r in REPORTS]
        self.assertEqual(len(slugs), len(set(slugs)))
        self.assertEqual(len(REPORT_INDEX), 9)


class BaseReportTestCase(TestCase):
    def setUp(self):
        ensure_role_groups()
        self.theme = IrpTheme.objects.create(
            code_name="A.B", title="Качество услуг", version=3
        )
        self.tfoms_user = Employee.objects.create_user(
            username="rep_op", password="GoodPass!1", org=81000
        )
        self.smo_user = Employee.objects.create_user(
            username="rep_smo", password="GoodPass!1", org=81001
        )
        self.tfoms_user.groups.add(Group.objects.get(name="ОП1"))
        self.smo_user.groups.add(Group.objects.get(name="СП1"))

    def _make(self, owner=None, irp_type=1, how=1, way=1, zh_d=None,
              date_close=None, result=None, text="Текст", line_one=None, pr_out=None):
        owner = owner or self.tfoms_user
        if date_close and result is None:
            result = 2
        return Irp.objects.create(
            n_irp=str(uuid.uuid4()),
            irp_type=irp_type,
            date_create=datetime.date.today(),
            way=way,
            how=how,
            theme=self.theme,
            otv_t=1,
            otv_kon=owner.org,
            employee_one=owner,
            data_plan=datetime.date.today() + datetime.timedelta(days=30),
            z_f="Петров",
            text=text,
            zh_d=zh_d,
            date_close=date_close,
            result=result,
            status=Irp.Status.CLOSED if date_close else Irp.Status.REGISTERED,
            line_one=line_one,
            pr_out=pr_out,
        )

    def _filters(self, **kwargs):
        return ReportFilters(**kwargs)


class ReportQueriesTests(BaseReportTestCase):
    def test_each_report_build_uses_single_query_without_per_row_fetches(self):
        self._make(irp_type=1, how=1, date_close=datetime.date.today())
        self._make(irp_type=2, how=1, zh_d="1.1")
        self._make(irp_type=4)

        for report in REPORTS:
            with self.subTest(report=report.slug), CaptureQueriesContext(connection) as queries:
                report.build(self.tfoms_user.org, self._filters())
            self.assertEqual(len(queries), 1)

    def test_scope_smo_sees_only_own(self):
        self._make(irp_type=2)
        self._make(owner=self.smo_user, irp_type=2)
        rows = REPORT_INDEX["r4_complaints"].build(self.smo_user.org, self._filters())
        self.assertEqual(rows[0]["total"], 1)

    def test_scope_uses_immutable_owner_not_mutable_responsible_org(self):
        own = self._make(owner=self.smo_user, irp_type=2)
        own.otv_kon = self.tfoms_user.org
        own.save(update_fields=["otv_kon"])
        foreign = self._make(owner=self.tfoms_user, irp_type=2)
        foreign.otv_kon = self.smo_user.org
        foreign.save(update_fields=["otv_kon"])

        for report in REPORTS:
            with self.subTest(report=report.slug):
                rows = report.build(self.smo_user.org, self._filters())
                if report.slug == "r4_complaints":
                    self.assertEqual(rows[-1]["total"], 1)
                elif report.slug == "r9_personal":
                    self.assertEqual([row["n_irp"] for row in rows], [own.n_irp])

    def test_tfoms_sees_all(self):
        self._make(irp_type=2)
        self._make(owner=self.smo_user, irp_type=2)
        rows = REPORT_INDEX["r4_complaints"].build(self.tfoms_user.org, self._filters())
        self.assertEqual(rows[0]["total"], 2)

    def test_only_closed_filter(self):
        self._make(irp_type=2, date_close=datetime.date.today())
        self._make(irp_type=2)
        rows = REPORT_INDEX["r4_complaints"].build(
            self.tfoms_user.org, self._filters(only_closed=True)
        )
        self.assertEqual(rows[0]["total"], 1)
        self.assertEqual(rows[0]["closed"], 1)

    def test_r1_by_volume_totals(self):
        self._make(irp_type=1, date_close=datetime.date.today())
        self._make(irp_type=2)
        rows = REPORT_INDEX["r1_volume"].build(self.tfoms_user.org, self._filters())
        total = rows[-1]
        self.assertEqual(total["total"], 2)
        self.assertEqual(total["t1"], 1)
        self.assertEqual(total["t2"], 1)
        self.assertEqual(total["closed"], 1)
        self.assertEqual(rows[0]["t3"], 0)
        self.assertEqual(rows[0]["t4"], 0)
        self.assertEqual(rows[0]["t5"], 0)

    def test_r2_by_type_percent(self):
        self._make(irp_type=1)
        self._make(irp_type=1)
        self._make(irp_type=2)
        rows = REPORT_INDEX["r2_type"].build(self.tfoms_user.org, self._filters())
        self.assertEqual(rows[-1]["total"], 3)
        self.assertEqual(rows[0]["percent"], "66,7")

    def test_r3_protection_buckets(self):
        self._make(irp_type=2, zh_d="1.1")
        self._make(irp_type=2, zh_d="1.2")
        self._make(irp_type=2, zh_d="2")
        rows = REPORT_INDEX["r3_protection"].build(self.tfoms_user.org, self._filters())
        detail = rows[0]
        self.assertEqual(detail["well"], 2)
        self.assertEqual(detail["pre"], 1)
        self.assertEqual(detail["court"], 1)
        self.assertEqual(detail["bad"], 1)
        self.assertEqual(detail["total"], 3)
        self.assertEqual(rows[-1]["total"], 3)

    def test_r5_applications_statuses(self):
        self._make(irp_type=4, result=3, date_close=datetime.date.today())
        self._make(irp_type=4, result=4, date_close=datetime.date.today())
        self._make(irp_type=4)
        rows = REPORT_INDEX["r5_applications"].build(self.tfoms_user.org, self._filters())
        total = rows[-1]
        self.assertEqual(total["total"], 3)
        self.assertEqual(total["satisfied"], 1)
        self.assertEqual(total["rejected"], 1)
        self.assertEqual(total["pending"], 1)

    def test_r6_clarification_hotline(self):
        self._make(irp_type=1, how=1, result=1, date_close=datetime.date.today())
        self._make(irp_type=1, how=3)
        rows = REPORT_INDEX["r6_clarification"].build(self.tfoms_user.org, self._filters())
        total = rows[-1]
        self.assertEqual(total["total"], 2)
        self.assertEqual(total["hotline"], 1)
        self.assertEqual(total["consulted"], 1)

    def test_r7_hotline_complaints_lines(self):
        self._make(irp_type=2, how=1, line_one=1)
        self._make(irp_type=2, how=1, line_one=2, pr_out=1)
        self._make(irp_type=2, how=2)  # не горячая линия
        rows = REPORT_INDEX["r7_hotline_complaints"].build(
            self.tfoms_user.org, self._filters()
        )
        total = rows[-1]
        self.assertEqual(total["total"], 2)
        self.assertEqual(total["op1"], 1)
        self.assertEqual(total["op2"], 1)
        self.assertEqual(total["redirected"], 1)

    def test_r9_personal_rows(self):
        irp = self._make(text="Многофункциональный центр", date_close=datetime.date.today(), result=2)
        rows = REPORT_INDEX["r9_personal"].build(self.tfoms_user.org, self._filters())
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["fio"], "Петров")
        self.assertEqual(rows[0]["status"], "закрыто")
        self.assertEqual(rows[0]["text"], "Многофункциональный центр")
        self.assertEqual(rows[0]["n_irp"], irp.n_irp)

    def test_date_range_filter(self):
        self._make()
        old = self._make()
        old.date_create = datetime.date(2020, 1, 1)
        old.save()
        rows = REPORT_INDEX["r2_type"].build(
            self.tfoms_user.org,
            self._filters(
                date_from=datetime.date.today() - datetime.timedelta(days=1),
                date_to=datetime.date.today(),
            ),
        )
        self.assertEqual(rows[-1]["total"], 1)


class FilterFormTests(TestCase):
    def setUp(self):
        self.tfoms_user = Employee.objects.create_user(
            username="ff_op", password="GoodPass!1", org=81000
        )
        self.smo_user = Employee.objects.create_user(
            username="ff_smo", password="GoodPass!1", org=81001
        )

    def test_invalid_period(self):
        form = ReportFilterForm(
            {
                "date_from": datetime.date(2026, 1, 10).isoformat(),
                "date_to": datetime.date(2026, 1, 1).isoformat(),
            },
            user=self.tfoms_user,
        )
        self.assertFalse(form.is_valid())

    def test_period_boundary_is_required_server_side(self):
        form = ReportFilterForm({"how": "1"}, user=self.tfoms_user)

        self.assertFalse(form.is_valid())
        self.assertIn(
            "Укажите хотя бы дату начала или окончания периода",
            str(form.non_field_errors()),
        )

    def test_to_filters(self):
        form = ReportFilterForm(
            {
                "date_from": datetime.date(2026, 1, 1).isoformat(),
                "how": "1",
                "only_closed": "on",
            },
            user=self.tfoms_user,
        )
        self.assertTrue(form.is_valid())
        f = form.to_filters()
        self.assertEqual(f.how, 1)
        self.assertTrue(f.only_closed)

    def test_smo_own_only_index(self):
        form = ReportFilterForm(user=self.smo_user)
        choices = dict(form.fields["otv_kon"].choices)
        self.assertEqual(list(choices), ["", 81001])


class ReportScreenTests(BaseReportTestCase):
    def test_index_requires_login(self):
        resp = self.client.get(reverse("reports:index"))
        self.assertEqual(resp.status_code, 302)

    def test_index_denies_user_without_role(self):
        user = Employee.objects.create_user(
            username="rep_no_role", password="GoodPass!1", org=81000
        )
        self.client.force_login(user)
        self.assertEqual(self.client.get(reverse("reports:index")).status_code, 403)

    def test_index_lists_reports(self):
        self.client.force_login(self.tfoms_user)
        resp = self.client.get(reverse("reports:index"))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Приложение №1")
        self.assertContains(resp, "Приложение №9")

    def test_detail_unknown_404(self):
        self.client.force_login(self.tfoms_user)
        resp = self.client.get(reverse("reports:detail", args=["nope"]))
        self.assertEqual(resp.status_code, 404)

    def test_detail_with_filters_renders_table(self):
        self._make(irp_type=2)
        self.client.force_login(self.tfoms_user)
        resp = self.client.get(
            reverse("reports:detail", args=["r4_complaints"]),
            {"how": "1", "date_from": datetime.date.today().isoformat()},
        )
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Качество услуг")
        self.assertContains(
            resp,
            'class="data data--responsive" data-client-sort data-table-key="reports-detail"',
        )
        self.assertContains(resp, 'class="responsive-row')
        self.assertContains(resp, 'data-label="Причина"')

    def test_export_xlsx(self):
        self._make(irp_type=2)
        self.client.force_login(self.tfoms_user)
        resp = self.client.get(
            reverse("reports:export", args=["r4_complaints", "xlsx"]),
            {"how": "1", "date_from": datetime.date.today().isoformat()},
        )
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("pril4", resp["Content-Disposition"])

    def test_export_pdf(self):
        self._make(irp_type=2)
        self.client.force_login(self.tfoms_user)
        resp = self.client.get(
            reverse("reports:export", args=["r4_complaints", "pdf"]),
            {"how": "1", "date_from": datetime.date.today().isoformat()},
        )
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp["Content-Type"], "application/pdf")
        self.assertTrue(resp.content.startswith(b"%PDF"))

    def test_smo_personal_export_cannot_cross_owner_scope_after_redirect(self):
        own = self._make(owner=self.smo_user)
        own.otv_kon = self.tfoms_user.org
        own.save(update_fields=["otv_kon"])
        foreign = self._make(owner=self.tfoms_user)
        foreign.otv_kon = self.smo_user.org
        foreign.save(update_fields=["otv_kon"])
        self.client.force_login(self.smo_user)

        response = self.client.get(
            reverse("reports:export", args=["r9_personal", "xlsx"]),
            {"date_from": datetime.date.today().isoformat()},
        )

        self.assertEqual(response.status_code, 200)
        worksheet = load_workbook(io.BytesIO(response.content)).active
        headers = [cell.value for cell in worksheet[1]]
        number_column = headers.index("№ обращения") + 1
        numbers = [
            worksheet.cell(row=row, column=number_column).value
            for row in range(2, worksheet.max_row + 1)
        ]
        self.assertEqual(numbers, [own.n_irp])
        self.assertNotIn(foreign.n_irp, numbers)

    def test_superuser_outside_tfoms_retains_global_report_scope(self):
        self._make(owner=self.tfoms_user, irp_type=2)
        self._make(owner=self.smo_user, irp_type=2)
        root = Employee.objects.create_superuser(
            username="report_foreign_org_root",
            password="GoodPass!1",
            org=81007,
        )
        self.client.force_login(root)

        response = self.client.get(
            reverse("reports:detail", args=["r4_complaints"]),
            {"date_from": datetime.date.today().isoformat()},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["rows"][-1]["total"], 2)
        self.assertEqual(
            set(dict(response.context["form"].fields["otv_kon"].choices)),
            {"", 81000, 81001, 81007, 81008},
        )
        exported = self.client.get(
            reverse("reports:export", args=["r4_complaints", "xlsx"]),
            {"date_from": datetime.date.today().isoformat()},
        )
        worksheet = load_workbook(io.BytesIO(exported.content)).active
        headers = [cell.value for cell in worksheet[1]]
        total_column = headers.index("Кол-во") + 1

        self.assertEqual(exported.status_code, 200)
        self.assertEqual(
            worksheet.cell(row=worksheet.max_row, column=total_column).value,
            2,
        )

    def test_export_invalid_fmt_404(self):
        self.client.force_login(self.tfoms_user)
        resp = self.client.get(
            reverse("reports:export", args=["r4_complaints", "doc"])
        )
        self.assertEqual(resp.status_code, 404)

    def test_export_invalid_filters_400(self):
        self.client.force_login(self.tfoms_user)
        resp = self.client.get(
            reverse("reports:export", args=["r4_complaints", "xlsx"]),
            {"date_from": "2026-02-01", "date_to": "2026-01-01"},
        )
        self.assertEqual(resp.status_code, 400)

    def test_direct_preview_and_export_without_period_are_rejected(self):
        self._make(irp_type=2)
        self.client.force_login(self.tfoms_user)

        preview = self.client.get(
            reverse("reports:detail", args=["r4_complaints"]), {"how": "1"}
        )
        export = self.client.get(
            reverse("reports:export", args=["r4_complaints", "xlsx"]),
            {"how": "1"},
        )

        self.assertEqual(preview.status_code, 200)
        self.assertContains(
            preview, "Укажите хотя бы дату начала или окончания периода"
        )
        self.assertContains(preview, 'data-key="report-filters" open')
        self.assertIsNone(preview.context["rows"])
        self.assertEqual(export.status_code, 400)


class ExportBytesTests(BaseReportTestCase):
    def test_write_xlsx_returns_zip(self):
        spec = REPORT_INDEX["r4_complaints"]
        rows = spec.build(self.tfoms_user.org, self._filters())
        data = write_xlsx_bytes(spec, rows)
        self.assertTrue(data.startswith(b"PK"))

    def test_xlsx_user_text_cannot_become_formula(self):
        spec = REPORT_INDEX["r4_complaints"]
        row = dict.fromkeys(spec.keys, "")
        row[spec.keys[0]] = '=HYPERLINK("https://example.invalid")'

        workbook = load_workbook(io.BytesIO(write_xlsx_bytes(spec, [row])))
        cell = workbook.active.cell(row=2, column=1)

        self.assertEqual(cell.data_type, "s")
        self.assertTrue(cell.value.startswith("'="))

    def test_pdf_escapes_user_markup(self):
        spec = REPORT_INDEX["r4_complaints"]
        row = dict.fromkeys(spec.keys, "")
        row[spec.keys[0]] = "<broken & text>"

        data = write_pdf(spec, [row])

        self.assertTrue(data.startswith(b"%PDF"))

    def test_write_pdf_returns_pdf(self):
        spec = REPORT_INDEX["r9_personal"]
        self._make(irp_type=1)
        rows = spec.build(self.tfoms_user.org, self._filters())
        data = write_pdf(spec, rows)
        self.assertTrue(data.startswith(b"%PDF"))
